"""Build the DCF valuation workbook (Assumptions, Segment_Revenue, Working_Capital, Capex tabs).

Run:  python dcf_model/build_dcf_model.py   ->   dcf_model/DCF_Model.xlsx

Layout conventions shared by every time-series tab (keep new tabs on these):
  * Column B = row labels
  * Column C = base year (last actual fiscal year, period t = 0)
  * Columns D:H = projection years 1-5 (periods t = 1..5)
"""
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).with_name("DCF_Model.xlsx")

FONT = "Arial"
BLUE, BLACK, GREEN, WHITE = "0000FF", "000000", "008000", "FFFFFF"
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
SECTION_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
TOP_BORDER = Border(top=Side(style="thin", color="000000"))

FMT_USD = '$#,##0.0_);($#,##0.0);"-"_)'
FMT_PCT = '0.0%_);(0.0%);"-"_)'
FMT_YEAR = "0"
FMT_DATE = "mmm d, yyyy"
FMT_DAYS = '0.0_);(0.0);"-"_)'
FMT_INTENSITY = '$0.00_);($0.00);"-"_)'

N_SEGMENTS = 8
YEAR_COLS = "CDEFGH"  # C = base year, D:H = projection years 1-5
PROJ_COLS = YEAR_COLS[1:]

# Example inputs so the model shows numbers out of the box -- replace with real data.
EXAMPLE_SEGMENTS = [
    # name, current revenue ($mm), growth / yr, operating margin
    ("Products", 1200.0, 0.06, 0.22),
    ("Services", 650.0, 0.10, 0.30),
    ("Licensing & Other", 150.0, 0.03, 0.08),
]

# Example working capital / capex drivers (base year + 5 projected years) -- replace.
EXAMPLE_COGS_PCT = [0.55] * 6
EXAMPLE_DSO = [45, 45, 44, 44, 43, 43]
EXAMPLE_DIO = [60, 60, 58, 57, 56, 55]
EXAMPLE_DPO = [40] * 6
EXAMPLE_BALANCES = (245.0, 180.0, 120.0)  # base-year AR, inventory, AP ($mm)
EXAMPLE_MAINT_CAPEX = 0.03
EXAMPLE_CAPEX_INTENSITY = [0.25, 0.15, 0.10]  # $ capex per $1 of incremental revenue

# ---- Assumptions tab cell map --------------------------------------------
A = "Assumptions"
A_COMPANY, A_TICKER, A_UNITS, A_BASE_YEAR, A_VAL_DATE = 7, 8, 9, 10, 11
A_TAX, A_WACC, A_TGR, A_TGR_CHECK = 14, 15, 16, 17
A_SEG_HDR = 20
A_SEG_FIRST = A_SEG_HDR + 1
A_SEG_LAST = A_SEG_FIRST + N_SEGMENTS - 1
A_SEG_TOTAL = A_SEG_LAST + 1
A_NOTES = A_SEG_TOTAL + 2  # notes occupy A_NOTES+1 .. A_NOTES+5

# Sections 4-5 (working capital & capex drivers), added below the Model Notes block
A_WC_HDR = A_NOTES + 7
A_WC_FY = A_WC_HDR + 1
A_COGS_PCT, A_DSO, A_DIO, A_DPO, A_DAYS = range(A_WC_FY + 1, A_WC_FY + 6)
A_BAL_HDR = A_DAYS + 1
A_AR0, A_INV0, A_AP0 = range(A_BAL_HDR + 1, A_BAL_HDR + 4)
A_CPX_HDR = A_AP0 + 2
A_MAINT = A_CPX_HDR + 1
A_INT_HDR = A_MAINT + 1
A_INT_FIRST = A_INT_HDR + 1
A_INT_LAST = A_INT_FIRST + N_SEGMENTS - 1
A_DRV_NOTES = A_INT_LAST + 2

# ---- Segment_Revenue tab row map ------------------------------------------
S = "Segment_Revenue"
S_FY, S_T = 4, 5
S_REV_HDR = 7
S_REV_FIRST = S_REV_HDR + 1
S_REV_TOTAL = S_REV_FIRST + N_SEGMENTS
S_REV_GROWTH = S_REV_TOTAL + 1
S_OI_HDR = S_REV_GROWTH + 2
S_OI_FIRST = S_OI_HDR + 1
S_OI_TOTAL = S_OI_FIRST + N_SEGMENTS
S_OI_MARGIN = S_OI_TOTAL + 1
S_SUM_HDR = S_OI_MARGIN + 2
S_SUM_REV, S_SUM_GROWTH, S_SUM_EBIT, S_SUM_MARGIN = range(S_SUM_HDR + 1, S_SUM_HDR + 5)
S_CHECK = S_SUM_MARGIN + 2

# ---- Working_Capital tab row map ------------------------------------------
W = "Working_Capital"
W_DRV_HDR = 7
W_REV, W_COGS_PCT, W_COGS, W_DSO, W_DIO, W_DPO = range(W_DRV_HDR + 1, W_DRV_HDR + 7)
W_BAL_HDR = W_DPO + 2
W_AR, W_INV, W_AP, W_NWC, W_NWC_PCT = range(W_BAL_HDR + 1, W_BAL_HDR + 6)
W_CHG_HDR = W_NWC_PCT + 2
W_CHG = W_CHG_HDR + 1

# ---- Capex tab row map ----------------------------------------------------
X = "Capex"
X_MAINT_HDR = 7
X_REV, X_MAINT_PCT, X_MAINT = range(X_MAINT_HDR + 1, X_MAINT_HDR + 4)
X_INC_HDR = X_MAINT + 2
X_INC_FIRST = X_INC_HDR + 1
X_INC_TOTAL = X_INC_FIRST + N_SEGMENTS
X_GC_HDR = X_INC_TOTAL + 2
X_GC_FIRST = X_GC_HDR + 1
X_GC_TOTAL = X_GC_FIRST + N_SEGMENTS
X_SUM_HDR = X_GC_TOTAL + 2
X_SUM_MAINT, X_SUM_GROWTH, X_SUM_TOTAL, X_SUM_PCT = range(X_SUM_HDR + 1, X_SUM_HDR + 5)


def put(ws, ref, value, color=BLACK, fmt=None, bold=False, italic=False, fill=None, size=10):
    c = ws[ref]
    c.value = value
    c.font = Font(name=FONT, color=color, bold=bold, italic=italic, size=size)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


def section(ws, row, text, first="B", last="F"):
    for col in range(ord(first), ord(last) + 1):
        ws[f"{chr(col)}{row}"].fill = SECTION_FILL
    put(ws, f"{first}{row}", text, color=WHITE, bold=True)


def total_style(ws, row, first="B", last="H"):
    for col in range(ord(first), ord(last) + 1):
        c = ws[f"{chr(col)}{row}"]
        c.border = TOP_BORDER
        c.font = Font(name=FONT, color=c.font.color, bold=True, size=10)


def add_name(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def pct_validation(ws, rng, lo, hi):
    dv = DataValidation(
        type="decimal", operator="between", formula1=str(lo), formula2=str(hi),
        allow_blank=True, showErrorMessage=True, errorTitle="Invalid percentage",
        error=f"Enter a percentage between {lo:.0%} and {hi:.0%} (e.g. 25% or 0.25).",
    )
    ws.add_data_validation(dv)
    dv.add(rng)


def num_validation(ws, rng, lo, hi, what):
    dv = DataValidation(
        type="decimal", operator="between", formula1=str(lo), formula2=str(hi),
        allow_blank=True, showErrorMessage=True, errorTitle=f"Invalid {what}",
        error=f"Enter a {what} between {lo} and {hi}.",
    )
    ws.add_data_validation(dv)
    dv.add(rng)


def seg_label(a_row):
    return f'=IF({A}!$B{a_row}="","",{A}!$B{a_row})'


def build_assumptions(wb):
    ws = wb.active
    ws.title = A
    ws.sheet_view.showGridLines = False
    for col, width in zip("ABCDEF", (2, 40, 16, 18, 18, 50)):
        ws.column_dimensions[col].width = width

    put(ws, "B1", "DCF Valuation Model - Assumptions & Config", bold=True, size=14)
    put(ws, "B2", "All currency figures in $ millions unless noted. "
                  "Edit only blue cells with yellow fill; everything else is formula-driven.", italic=True)
    put(ws, "B4", "Legend:", bold=True)
    put(ws, "C4", "Input", color=BLUE, fill=INPUT_FILL)
    put(ws, "D4", "Formula", color=BLACK)
    put(ws, "E4", "Link from other tab", color=GREEN)

    # 1. General / config
    section(ws, 6, "1. General / Config")
    rows = [
        (A_COMPANY, "Company Name", "Example Co.", None, "Example placeholder - replace"),
        (A_TICKER, "Ticker", "EXCO", None, "Example placeholder - replace"),
        (A_UNITS, "Units", "$ in millions", None, "Label only; shown in tab headers"),
        (A_BASE_YEAR, "Last Actual Fiscal Year (Base Year)", 2025, FMT_YEAR,
         "Projections run Base Year +1 through +5"),
        (A_VAL_DATE, "Valuation Date", date(2026, 10, 13), FMT_DATE, "Example placeholder - replace"),
    ]
    for r, label, val, fmt, note in rows:
        put(ws, f"B{r}", label)
        put(ws, f"C{r}", val, color=BLUE, fmt=fmt, fill=INPUT_FILL)
        put(ws, f"D{r}", note, italic=True)

    # 2. Company-wide valuation inputs
    section(ws, 13, "2. Company-Wide Valuation Inputs")
    rows = [
        (A_TAX, "Tax Rate", 0.25, "Example placeholder - effective/marginal cash tax rate"),
        (A_WACC, "Discount Rate (WACC)", 0.09, "Example placeholder - replace with WACC build"),
        (A_TGR, "Terminal Growth Rate", 0.025, "Example placeholder - long-run nominal growth"),
    ]
    for r, label, val, note in rows:
        put(ws, f"B{r}", label)
        put(ws, f"C{r}", val, color=BLUE, fmt=FMT_PCT, fill=INPUT_FILL)
        put(ws, f"D{r}", note, italic=True)
    put(ws, f"B{A_TGR_CHECK}", "Check: Terminal Growth < WACC", italic=True)
    put(ws, f"C{A_TGR_CHECK}", f'=IF(C{A_TGR}<C{A_WACC},"OK","ERROR")', bold=True)
    pct_validation(ws, f"C{A_TAX}:C{A_WACC}", 0, 1)
    pct_validation(ws, f"C{A_TGR}", -0.1, 0.2)

    # 3. Segment inputs
    section(ws, A_SEG_HDR - 1, f"3. Segment Inputs (up to {N_SEGMENTS} segments)")
    headers = ["Segment Name", "Current Revenue ($mm)", "Projected Growth Rate (% / yr)",
               "Operating Margin (%)", "Notes / Source"]
    for col, text in zip("BCDEF", headers):
        c = put(ws, f"{col}{A_SEG_HDR}", text, bold=True, fill=HEADER_FILL)
        c.alignment = Alignment(wrap_text=True, horizontal="left" if col in "BF" else "center",
                                vertical="center")
    ws.row_dimensions[A_SEG_HDR].height = 30

    for i in range(N_SEGMENTS):
        r = A_SEG_FIRST + i
        name, rev, growth, margin = EXAMPLE_SEGMENTS[i] if i < len(EXAMPLE_SEGMENTS) else (None,) * 4
        put(ws, f"B{r}", name, color=BLUE, fill=INPUT_FILL)
        put(ws, f"C{r}", rev, color=BLUE, fmt=FMT_USD, fill=INPUT_FILL)
        put(ws, f"D{r}", growth, color=BLUE, fmt=FMT_PCT, fill=INPUT_FILL)
        put(ws, f"E{r}", margin, color=BLUE, fmt=FMT_PCT, fill=INPUT_FILL)
        put(ws, f"F{r}", "Example placeholder - replace with company data" if name else None,
            color=BLUE, italic=True)
    pct_validation(ws, f"D{A_SEG_FIRST}:D{A_SEG_LAST}", -1, 1)
    pct_validation(ws, f"E{A_SEG_FIRST}:E{A_SEG_LAST}", -1, 1)

    f, l, t = A_SEG_FIRST, A_SEG_LAST, A_SEG_TOTAL
    put(ws, f"B{t}", "Total / Revenue-Weighted Average")
    put(ws, f"C{t}", f"=SUM(C{f}:C{l})", fmt=FMT_USD)
    put(ws, f"D{t}", f"=IF(C{t}=0,0,SUMPRODUCT(C{f}:C{l},D{f}:D{l})/C{t})", fmt=FMT_PCT)
    put(ws, f"E{t}", f"=IF(C{t}=0,0,SUMPRODUCT(C{f}:C{l},E{f}:E{l})/C{t})", fmt=FMT_PCT)
    total_style(ws, t, last="F")

    # Notes
    put(ws, f"B{A_NOTES}", "Model Notes", bold=True)
    notes = [
        f"To add a segment, fill the next empty slot (rows {A_SEG_FIRST}-{A_SEG_LAST}); "
        f"{S} picks it up automatically. Blank slots calculate as zero.",
        "Growth rate and operating margin are applied uniformly to all 5 projection years; "
        "base-year operating income = current revenue x operating margin.",
        f"Time-series tabs share one column layout: C = base year (actual), D:H = years 1-5. "
        f"Link downstream tabs to the Total Company Summary block on {S}.",
        "Named ranges: CompanyName, Ticker, Units, BaseYear, ValuationDate, TaxRate, WACC, "
        "TerminalGrowth, Seg_Names, Seg_Revenue, Seg_Growth, Seg_Margin, FY_Labels, Period, "
        "Total_Revenue, Total_EBIT.",
        "Add new driver sections (e.g. 4. Working Capital, 5. Capex & D&A) below this notes block.",
    ]
    for i, text in enumerate(notes, start=1):
        put(ws, f"B{A_NOTES + i}", f"{i}. {text}", italic=True)

    for name, row in [("CompanyName", A_COMPANY), ("Ticker", A_TICKER), ("Units", A_UNITS),
                      ("BaseYear", A_BASE_YEAR), ("ValuationDate", A_VAL_DATE),
                      ("TaxRate", A_TAX), ("WACC", A_WACC), ("TerminalGrowth", A_TGR)]:
        add_name(wb, name, f"{A}!$C${row}")
    for name, col in [("Seg_Names", "B"), ("Seg_Revenue", "C"), ("Seg_Growth", "D"), ("Seg_Margin", "E")]:
        add_name(wb, name, f"{A}!${col}${A_SEG_FIRST}:${col}${A_SEG_LAST}")


def build_segment_revenue(wb):
    ws = wb.create_sheet(S)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 38
    for col in YEAR_COLS + "I":
        ws.column_dimensions[col].width = 13

    put(ws, "B1", '=CompanyName&" - Segment Revenue & Operating Income Build"', bold=True, size=14)
    put(ws, "B2", '="("&Units&"; all drivers link to the Assumptions tab)"', italic=True)

    # Timeline header
    put(ws, f"B{S_FY}", "Fiscal Year", bold=True)
    put(ws, f"C{S_FY}", '="FY"&BaseYear&"A"', bold=True)
    for col in PROJ_COLS:
        put(ws, f"{col}{S_FY}", f'="FY"&(BaseYear+{col}${S_T})&"E"', bold=True)
    put(ws, f"I{S_FY}", "5-Yr CAGR", bold=True)
    put(ws, f"B{S_T}", "Projection Period (t)", italic=True)
    put(ws, f"C{S_T}", 0, color=BLUE, italic=True)
    for prev, col in zip(YEAR_COLS, PROJ_COLS):
        put(ws, f"{col}{S_T}", f"={prev}{S_T}+1", italic=True)
    for col in "BCDEFGHI":
        ws[f"{col}{S_FY}"].fill = HEADER_FILL
        ws[f"{col}{S_FY}"].alignment = Alignment(horizontal="left" if col == "B" else "right")
        ws[f"{col}{S_T}"].alignment = Alignment(horizontal="left" if col == "B" else "right")

    def cagr(row):
        return f"=IF(C{row}<=0,0,(H{row}/C{row})^(1/$H${S_T})-1)"

    # Revenue by segment
    section(ws, S_REV_HDR, "Revenue by Segment", last="I")
    for i in range(N_SEGMENTS):
        r, a = S_REV_FIRST + i, A_SEG_FIRST + i
        put(ws, f"B{r}", seg_label(a), color=GREEN)
        put(ws, f"C{r}", f"={A}!$C{a}", color=GREEN, fmt=FMT_USD)
        for prev, col in zip(YEAR_COLS, PROJ_COLS):
            put(ws, f"{col}{r}", f"={prev}{r}*(1+{A}!$D{a})", fmt=FMT_USD)
        put(ws, f"I{r}", cagr(r), fmt=FMT_PCT)
    put(ws, f"B{S_REV_TOTAL}", "Total Revenue")
    for col in YEAR_COLS:
        put(ws, f"{col}{S_REV_TOTAL}", f"=SUM({col}{S_REV_FIRST}:{col}{S_REV_TOTAL - 1})", fmt=FMT_USD)
    put(ws, f"I{S_REV_TOTAL}", cagr(S_REV_TOTAL), fmt=FMT_PCT)
    total_style(ws, S_REV_TOTAL, last="I")
    put(ws, f"B{S_REV_GROWTH}", "  % Growth", italic=True)
    for prev, col in zip(YEAR_COLS, PROJ_COLS):
        put(ws, f"{col}{S_REV_GROWTH}",
            f"=IF({prev}{S_REV_TOTAL}=0,0,{col}{S_REV_TOTAL}/{prev}{S_REV_TOTAL}-1)",
            fmt=FMT_PCT, italic=True)

    # Operating income by segment
    section(ws, S_OI_HDR, "Operating Income by Segment", last="I")
    for i in range(N_SEGMENTS):
        r, a, rev = S_OI_FIRST + i, A_SEG_FIRST + i, S_REV_FIRST + i
        put(ws, f"B{r}", seg_label(a), color=GREEN)
        for col in YEAR_COLS:
            put(ws, f"{col}{r}", f"={col}{rev}*{A}!$E{a}", fmt=FMT_USD)
        put(ws, f"I{r}", cagr(r), fmt=FMT_PCT)
    put(ws, f"B{S_OI_TOTAL}", "Total Operating Income")
    for col in YEAR_COLS:
        put(ws, f"{col}{S_OI_TOTAL}", f"=SUM({col}{S_OI_FIRST}:{col}{S_OI_TOTAL - 1})", fmt=FMT_USD)
    put(ws, f"I{S_OI_TOTAL}", cagr(S_OI_TOTAL), fmt=FMT_PCT)
    total_style(ws, S_OI_TOTAL, last="I")
    put(ws, f"B{S_OI_MARGIN}", "  % Operating Margin", italic=True)
    for col in YEAR_COLS:
        put(ws, f"{col}{S_OI_MARGIN}",
            f"=IF({col}{S_REV_TOTAL}=0,0,{col}{S_OI_TOTAL}/{col}{S_REV_TOTAL})", fmt=FMT_PCT, italic=True)

    # Total company summary -- the link point for downstream tabs
    section(ws, S_SUM_HDR, "Total Company Summary (link downstream tabs to these rows)", last="I")
    put(ws, f"B{S_SUM_REV}", "Total Revenue", bold=True)
    put(ws, f"B{S_SUM_GROWTH}", "  % Growth", italic=True)
    put(ws, f"B{S_SUM_EBIT}", "Total Operating Income (EBIT)", bold=True)
    put(ws, f"B{S_SUM_MARGIN}", "  % Operating Margin", italic=True)
    for col in YEAR_COLS:
        put(ws, f"{col}{S_SUM_REV}", f"={col}{S_REV_TOTAL}", fmt=FMT_USD, bold=True)
        put(ws, f"{col}{S_SUM_EBIT}", f"={col}{S_OI_TOTAL}", fmt=FMT_USD, bold=True)
        put(ws, f"{col}{S_SUM_MARGIN}", f"={col}{S_OI_MARGIN}", fmt=FMT_PCT, italic=True)
    for col in PROJ_COLS:
        put(ws, f"{col}{S_SUM_GROWTH}", f"={col}{S_REV_GROWTH}", fmt=FMT_PCT, italic=True)
    put(ws, f"I{S_SUM_REV}", cagr(S_SUM_REV), fmt=FMT_PCT, bold=True)
    put(ws, f"I{S_SUM_EBIT}", cagr(S_SUM_EBIT), fmt=FMT_PCT, bold=True)

    put(ws, f"B{S_CHECK}", "Check: base-year revenue ties to Assumptions", italic=True)
    put(ws, f"C{S_CHECK}", f'=IF(ABS(C{S_SUM_REV}-{A}!$C${A_SEG_TOTAL})<0.001,"OK","ERROR")', bold=True)

    ws.freeze_panes = f"C{S_T + 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    add_name(wb, "FY_Labels", f"{S}!$C${S_FY}:$H${S_FY}")
    add_name(wb, "Period", f"{S}!$C${S_T}:$H${S_T}")
    add_name(wb, "Total_Revenue", f"{S}!$C${S_SUM_REV}:$H${S_SUM_REV}")
    add_name(wb, "Total_EBIT", f"{S}!$C${S_SUM_EBIT}:$H${S_SUM_EBIT}")


def add_driver_inputs(wb):
    """Sections 4-5 on Assumptions: new input cells only, below the existing Model Notes."""
    ws = wb[A]
    for col, width in zip("GHI", (18, 18, 48)):
        ws.column_dimensions[col].width = width

    # 4. Working capital drivers -- year columns follow the C:H convention
    section(ws, A_WC_HDR, "4. Working Capital Drivers", last="I")
    put(ws, f"B{A_WC_FY}", "Fiscal Year", bold=True, fill=HEADER_FILL)
    for col in YEAR_COLS:
        c = put(ws, f"{col}{A_WC_FY}", f"={S}!{col}${S_FY}", color=GREEN, bold=True, fill=HEADER_FILL)
        c.alignment = Alignment(horizontal="right")
    put(ws, f"I{A_WC_FY}", "Notes / Source", bold=True, fill=HEADER_FILL)
    rows = [
        (A_COGS_PCT, "COGS (% of Revenue)", EXAMPLE_COGS_PCT, FMT_PCT,
         "Drives COGS for DIO / DPO; example placeholder"),
        (A_DSO, "DSO - Days Sales Outstanding", EXAMPLE_DSO, FMT_DAYS,
         "Receivables = DSO x Revenue / Days per Year"),
        (A_DIO, "DIO - Days Inventory Outstanding", EXAMPLE_DIO, FMT_DAYS,
         "Inventory = DIO x COGS / Days per Year"),
        (A_DPO, "DPO - Days Payables Outstanding", EXAMPLE_DPO, FMT_DAYS,
         "Payables = DPO x COGS / Days per Year"),
    ]
    for r, label, vals, fmt, note in rows:
        put(ws, f"B{r}", label)
        for col, v in zip(YEAR_COLS, vals):
            put(ws, f"{col}{r}", v, color=BLUE, fmt=fmt, fill=INPUT_FILL)
        put(ws, f"I{r}", note, italic=True)
    put(ws, f"B{A_DAYS}", "Days per Year")
    put(ws, f"C{A_DAYS}", 365, color=BLUE, fmt="0", fill=INPUT_FILL)
    put(ws, f"I{A_DAYS}", "Day-count basis for DSO / DIO / DPO", italic=True)
    pct_validation(ws, f"C{A_COGS_PCT}:H{A_COGS_PCT}", 0, 1)
    num_validation(ws, f"C{A_DSO}:H{A_DPO}", 0, 365, "number of days")

    put(ws, f"B{A_BAL_HDR}", "Base-Year Balances ($mm)", bold=True)
    put(ws, f"C{A_BAL_HDR}", f"={S}!C${S_FY}", color=GREEN, bold=True)
    ws[f"C{A_BAL_HDR}"].alignment = Alignment(horizontal="right")
    for r, label, val in zip((A_AR0, A_INV0, A_AP0),
                             ("Accounts Receivable", "Inventory", "Accounts Payable"), EXAMPLE_BALANCES):
        put(ws, f"B{r}", label)
        put(ws, f"C{r}", val, color=BLUE, fmt=FMT_USD, fill=INPUT_FILL)
        put(ws, f"I{r}", "Actual balance sheet value; example placeholder - replace", italic=True)

    # 5. Capex drivers
    section(ws, A_CPX_HDR, "5. Capex Drivers", last="I")
    put(ws, f"B{A_MAINT}", "Maintenance Capex (% of Total Revenue)")
    put(ws, f"C{A_MAINT}", EXAMPLE_MAINT_CAPEX, color=BLUE, fmt=FMT_PCT, fill=INPUT_FILL)
    put(ws, f"I{A_MAINT}", "Applied to total revenue each projection year; example placeholder", italic=True)
    pct_validation(ws, f"C{A_MAINT}", 0, 1)

    for col, text in (("B", "Segment"), ("C", "Growth Capex Intensity ($ per $1 incr. revenue)"),
                      ("I", "Notes / Source")):
        c = put(ws, f"{col}{A_INT_HDR}", text, bold=True, fill=HEADER_FILL)
        c.alignment = Alignment(wrap_text=True, horizontal="center" if col == "C" else "left",
                                vertical="center")
    ws.row_dimensions[A_INT_HDR].height = 42
    for i in range(N_SEGMENTS):
        r, a = A_INT_FIRST + i, A_SEG_FIRST + i
        put(ws, f"B{r}", f'=IF($B{a}="","",$B{a})')
        val = EXAMPLE_CAPEX_INTENSITY[i] if i < len(EXAMPLE_CAPEX_INTENSITY) else None
        put(ws, f"C{r}", val, color=BLUE, fmt=FMT_INTENSITY, fill=INPUT_FILL)
        put(ws, f"I{r}", "Example placeholder - replace" if val is not None else None, color=BLUE, italic=True)
    num_validation(ws, f"C{A_INT_FIRST}:C{A_INT_LAST}", 0, 10, "capex intensity")

    put(ws, f"B{A_DRV_NOTES}", "Driver Notes (sections 4-5)", bold=True)
    notes = [
        "Base-year days (column C) are for reference; base-year NWC uses the actual balances above.",
        "Growth capex = intensity x MAX(0, segment revenue increase); a shrinking segment adds no growth capex.",
        "Sign convention: Change in NWC and Total Capex are positive when they use cash - "
        "subtract both on the FCF tab.",
        "Named ranges: COGS_Pct, DSO_Days, DIO_Days, DPO_Days, DaysPerYear, MaintCapexPct, "
        "Seg_CapexIntensity, NWC, Change_NWC, Maint_Capex, Growth_Capex, Total_Capex.",
    ]
    for i, text in enumerate(notes, start=1):
        put(ws, f"B{A_DRV_NOTES + i}", f"{i}. {text}", italic=True)

    for name, row in [("COGS_Pct", A_COGS_PCT), ("DSO_Days", A_DSO), ("DIO_Days", A_DIO), ("DPO_Days", A_DPO)]:
        add_name(wb, name, f"{A}!$C${row}:$H${row}")
    add_name(wb, "DaysPerYear", f"{A}!$C${A_DAYS}")
    add_name(wb, "MaintCapexPct", f"{A}!$C${A_MAINT}")
    add_name(wb, "Seg_CapexIntensity", f"{A}!$C${A_INT_FIRST}:$C${A_INT_LAST}")


def timeseries_sheet(wb, name, title):
    """New tab with the shared header: title, units, FY labels and periods linked from Segment_Revenue."""
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 44
    for col in YEAR_COLS:
        ws.column_dimensions[col].width = 13
    put(ws, "B1", f'=CompanyName&" - {title}"', bold=True, size=14)
    put(ws, "B2", '="("&Units&"; drivers link to the Assumptions and Segment_Revenue tabs)"', italic=True)
    put(ws, f"B{S_FY}", "Fiscal Year", bold=True, fill=HEADER_FILL)
    put(ws, f"B{S_T}", "Projection Period (t)", italic=True)
    for col in YEAR_COLS:
        put(ws, f"{col}{S_FY}", f"={S}!{col}${S_FY}", color=GREEN, bold=True, fill=HEADER_FILL)
        put(ws, f"{col}{S_T}", f"={S}!{col}${S_T}", color=GREEN, italic=True)
        for r in (S_FY, S_T):
            ws[f"{col}{r}"].alignment = Alignment(horizontal="right")
    ws.freeze_panes = f"C{S_T + 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def build_working_capital(wb):
    ws = timeseries_sheet(wb, W, "Working Capital Schedule")

    section(ws, W_DRV_HDR, "Drivers", last="H")
    put(ws, f"B{W_REV}", "Total Revenue")
    put(ws, f"B{W_COGS_PCT}", "  COGS (% of Revenue)", italic=True)
    put(ws, f"B{W_COGS}", "COGS")
    put(ws, f"B{W_DSO}", "DSO (days)")
    put(ws, f"B{W_DIO}", "DIO (days)")
    put(ws, f"B{W_DPO}", "DPO (days)")
    for col in YEAR_COLS:
        put(ws, f"{col}{W_REV}", f"={S}!{col}{S_SUM_REV}", color=GREEN, fmt=FMT_USD)
        put(ws, f"{col}{W_COGS_PCT}", f"={A}!{col}{A_COGS_PCT}", color=GREEN, fmt=FMT_PCT, italic=True)
        put(ws, f"{col}{W_COGS}", f"={col}{W_REV}*{col}{W_COGS_PCT}", fmt=FMT_USD)
        for r, a in ((W_DSO, A_DSO), (W_DIO, A_DIO), (W_DPO, A_DPO)):
            put(ws, f"{col}{r}", f"={A}!{col}{a}", color=GREEN, fmt=FMT_DAYS)

    section(ws, W_BAL_HDR, "Working Capital Balances", last="H")
    days = f"{A}!$C${A_DAYS}"
    for r, label, a0, drv, base in ((W_AR, "Accounts Receivable", A_AR0, W_DSO, W_REV),
                                    (W_INV, "Inventory", A_INV0, W_DIO, W_COGS),
                                    (W_AP, "Accounts Payable", A_AP0, W_DPO, W_COGS)):
        put(ws, f"B{r}", label)
        put(ws, f"C{r}", f"={A}!$C${a0}", color=GREEN, fmt=FMT_USD)
        for col in PROJ_COLS:
            put(ws, f"{col}{r}", f"={col}{drv}*{col}{base}/{days}", fmt=FMT_USD)
    put(ws, f"B{W_NWC}", "Net Working Capital (AR + Inventory - AP)")
    put(ws, f"B{W_NWC_PCT}", "  NWC (% of Revenue)", italic=True)
    for col in YEAR_COLS:
        put(ws, f"{col}{W_NWC}", f"={col}{W_AR}+{col}{W_INV}-{col}{W_AP}", fmt=FMT_USD)
        put(ws, f"{col}{W_NWC_PCT}", f"=IF({col}{W_REV}=0,0,{col}{W_NWC}/{col}{W_REV})", fmt=FMT_PCT, italic=True)
    total_style(ws, W_NWC)

    section(ws, W_CHG_HDR, "Change in Net Working Capital (link FCF tab to this row)", last="H")
    put(ws, f"B{W_CHG}", "Change in NWC (increase = cash outflow)")
    for prev, col in zip(YEAR_COLS, PROJ_COLS):
        put(ws, f"{col}{W_CHG}", f"={col}{W_NWC}-{prev}{W_NWC}", fmt=FMT_USD)
    total_style(ws, W_CHG)

    add_name(wb, "NWC", f"{W}!$C${W_NWC}:$H${W_NWC}")
    add_name(wb, "Change_NWC", f"{W}!$C${W_CHG}:$H${W_CHG}")


def build_capex(wb):
    ws = timeseries_sheet(wb, X, "Capital Expenditure Schedule")

    section(ws, X_MAINT_HDR, "Maintenance Capex", last="H")
    put(ws, f"B{X_REV}", "Total Revenue")
    put(ws, f"B{X_MAINT_PCT}", "  Maintenance Capex (% of Revenue)", italic=True)
    put(ws, f"B{X_MAINT}", "Maintenance Capex")
    for col in YEAR_COLS:
        put(ws, f"{col}{X_REV}", f"={S}!{col}{S_SUM_REV}", color=GREEN, fmt=FMT_USD)
    for col in PROJ_COLS:
        put(ws, f"{col}{X_MAINT_PCT}", f"={A}!$C${A_MAINT}", color=GREEN, fmt=FMT_PCT, italic=True)
        put(ws, f"{col}{X_MAINT}", f"={col}{X_REV}*{col}{X_MAINT_PCT}", fmt=FMT_USD)

    section(ws, X_INC_HDR, "Incremental Revenue by Segment (year-over-year increase)", last="H")
    section(ws, X_GC_HDR, "Growth Capex by Segment (intensity x MAX(0, incremental revenue))", last="H")
    for i in range(N_SEGMENTS):
        inc, gc, a, sr = X_INC_FIRST + i, X_GC_FIRST + i, A_SEG_FIRST + i, S_REV_FIRST + i
        put(ws, f"B{inc}", seg_label(a), color=GREEN)
        put(ws, f"B{gc}", seg_label(a), color=GREEN)
        for prev, col in zip(YEAR_COLS, PROJ_COLS):
            put(ws, f"{col}{inc}", f"={S}!{col}{sr}-{S}!{prev}{sr}", fmt=FMT_USD)
            put(ws, f"{col}{gc}", f"=MAX(0,{col}{inc})*{A}!$C${A_INT_FIRST + i}", fmt=FMT_USD)
    put(ws, f"B{X_INC_TOTAL}", "Total Incremental Revenue")
    put(ws, f"B{X_GC_TOTAL}", "Total Growth Capex")
    for col in PROJ_COLS:
        put(ws, f"{col}{X_INC_TOTAL}", f"=SUM({col}{X_INC_FIRST}:{col}{X_INC_TOTAL - 1})", fmt=FMT_USD)
        put(ws, f"{col}{X_GC_TOTAL}", f"=SUM({col}{X_GC_FIRST}:{col}{X_GC_TOTAL - 1})", fmt=FMT_USD)
    total_style(ws, X_INC_TOTAL)
    total_style(ws, X_GC_TOTAL)

    section(ws, X_SUM_HDR, "Total Capex (link FCF tab to the Total Capex row)", last="H")
    put(ws, f"B{X_SUM_MAINT}", "Maintenance Capex")
    put(ws, f"B{X_SUM_GROWTH}", "Growth Capex")
    put(ws, f"B{X_SUM_TOTAL}", "Total Capex (positive = cash outflow)")
    put(ws, f"B{X_SUM_PCT}", "  Total Capex (% of Revenue)", italic=True)
    for col in PROJ_COLS:
        put(ws, f"{col}{X_SUM_MAINT}", f"={col}{X_MAINT}", fmt=FMT_USD)
        put(ws, f"{col}{X_SUM_GROWTH}", f"={col}{X_GC_TOTAL}", fmt=FMT_USD)
        put(ws, f"{col}{X_SUM_TOTAL}", f"={col}{X_SUM_MAINT}+{col}{X_SUM_GROWTH}", fmt=FMT_USD)
        put(ws, f"{col}{X_SUM_PCT}", f"=IF({col}{X_REV}=0,0,{col}{X_SUM_TOTAL}/{col}{X_REV})",
            fmt=FMT_PCT, italic=True)
    total_style(ws, X_SUM_TOTAL)

    add_name(wb, "Maint_Capex", f"{X}!$C${X_SUM_MAINT}:$H${X_SUM_MAINT}")
    add_name(wb, "Growth_Capex", f"{X}!$C${X_SUM_GROWTH}:$H${X_SUM_GROWTH}")
    add_name(wb, "Total_Capex", f"{X}!$C${X_SUM_TOTAL}:$H${X_SUM_TOTAL}")


def main():
    wb = Workbook()
    build_assumptions(wb)
    build_segment_revenue(wb)
    add_driver_inputs(wb)
    build_working_capital(wb)
    build_capex(wb)
    wb.save(OUT)
    print(f"Saved {OUT}")


if __name__ == "__main__":
    main()
